from __future__ import annotations

import asyncio
import json
import logging
from pathlib import Path
import re
from typing import Any

from docx import Document
import numpy as np

from agents.base import AgentFactory
from agents.base import BaseAgent
from agents.base import RESUME_STRATEGIES
from core.config import get_settings
from core.queue import JD_FLAGGED
from core.queue import JD_READY
from core.queue import queue_manager
from core.resume_store import get_resume_text as _get_resume_text


logger = logging.getLogger(__name__)
settings = get_settings()

RESUME_DIR = Path("/tmp/resumes")
RESUME_DIR.mkdir(parents=True, exist_ok=True)
BASE_RESUME_PATH = Path("/app/data/base_resume.docx")


def cosine_similarity(a: list[float], b: list[float]) -> float:
    a, b = np.array(a), np.array(b)
    denom = np.linalg.norm(a) * np.linalg.norm(b)
    return float(np.dot(a, b) / denom) if denom != 0 else 0.0


@AgentFactory.register("resume")
class ResumeEditorAgent(BaseAgent):
    agent_name = "resume_editor"

    def __init__(
        self,
        strategy_name: str = "keyword_injection",
        llm: Any | None = None,
        scorer: Any | None = None,
    ) -> None:
        super().__init__(llm=llm, scorer=scorer)
        self.strategy = RESUME_STRATEGIES.get(
            strategy_name,
            RESUME_STRATEGIES["keyword_injection"],
        )
        self.strategy_name = str(
            getattr(self.strategy, "strategy_name", "keyword_injection")
        )

    def validate_input(self, payload: Any) -> bool:
        if not isinstance(payload, dict):
            return False

        job_id = payload.get("job_id")
        company = payload.get("company_name") or payload.get("company")
        role = payload.get("role_title") or payload.get("role")
        jd_text = payload.get("jd_text") or payload.get("raw_text")
        return bool(job_id and company and role and jd_text)

    async def process(self, payload: Any) -> dict[str, Any]:
        payload_dict = dict(payload)
        job_id = str(payload_dict["job_id"])
        company = str(payload_dict.get("company_name") or payload_dict.get("company") or "")
        role = str(payload_dict.get("role_title") or payload_dict.get("role") or "")
        jd_text = str(payload_dict.get("jd_text") or payload_dict.get("raw_text") or "")
        synthesis = payload_dict.get("synthesis")
        required_skills = self._normalize_required_skills(
            payload_dict.get("required_skills_extended")
            or payload_dict.get("required_skills")
            or (synthesis or {}).get("key_skills_extended")
        )

        resume_text = self._load_resume_text()
        resume_embedding, jd_embedding = await asyncio.gather(
            self.llm.embed(resume_text),
            self.llm.embed(jd_text),
        )
        match_score_before = cosine_similarity(resume_embedding, jd_embedding)

        gap_analysis = await self._identify_gaps(resume_text, required_skills, jd_text)
        present_skills = self._normalize_required_skills(gap_analysis.get("present_skills"))
        gap_skills = self._normalize_required_skills(gap_analysis.get("missing_skills"))

        edit_result = await self.strategy.edit(resume_text, jd_text, gap_skills)
        match_score_after = await self._estimate_improved_score(
            match_score_before,
            edit_result,
            gap_skills,
            self.strategy_name,
        )

        edited_resume_file = RESUME_DIR / f"{job_id}_{self.strategy_name}.docx"
        edited_resume_path = self._save_edited_resume(
            edit_result,
            edited_resume_file,
            resume_text,
        )

        needs_human_review = match_score_after < settings.resume_match_threshold
        flag_reason = (
            self._build_flag_reason(match_score_after)
            if needs_human_review
            else None
        )
        changes_made = self._normalize_string_list(edit_result.get("changes"))

        result = dict(payload_dict)
        result.update(
            {
                "job_id": job_id,
                "company_name": company,
                "role_title": role,
                "jd_text": jd_text,
                "synthesis": synthesis,
                "required_skills_extended": required_skills,
                "resume_text": resume_text,
                "match_score_before": match_score_before,
                "match_score_after": match_score_after,
                "present_skills": present_skills,
                "gap_skills": gap_skills,
                "strategy_used": self.strategy_name,
                "edit_result": edit_result,
                "edited_resume_path": edited_resume_path,
                "needs_human_review": needs_human_review,
                "flag_reason": flag_reason,
                "changes_made": changes_made,
                "original_filename": BASE_RESUME_PATH.name,
                "edited_filename": edited_resume_file.name,
            }
        )
        return result

    async def emit_result(
        self,
        result: dict[str, Any],
        original_payload: Any,
    ) -> None:
        # Persist to DB before queuing
        try:
            from db.persistence import upsert_resume_version, upsert_application
            resume_version_id = await upsert_resume_version(result)
            await upsert_application(result, resume_version_id=resume_version_id or None)
        except Exception as exc:
            logger.exception("Failed to persist resume results to DB: %s", exc)

        if result.get("needs_human_review"):
            await queue_manager.publish(JD_FLAGGED, result, "jd.flagged.new")
            return

        await queue_manager.publish(JD_READY, result, "jd.ready.new")

    def _load_resume_text(self) -> str:
        """Load resume text via the central resume_store module."""
        return _get_resume_text()

    async def _identify_gaps(
        self,
        resume_text: str,
        required_skills: list[str],
        jd_text: str,
    ) -> dict[str, list[str]]:
        normalized_skills = self._normalize_required_skills(required_skills)
        if not normalized_skills:
            return {"missing_skills": [], "present_skills": []}

        prompt = (
            "Compare the resume against the job requirements and return JSON only.\n"
            'Return {"missing_skills": [], "present_skills": []}.\n\n'
            f"Required skills: {json.dumps(normalized_skills, ensure_ascii=False)}\n\n"
            f"Resume text:\n{resume_text}\n\n"
            f"Job description:\n{jd_text}"
        )

        try:
            response = await self.llm.complete(
                prompt=prompt,
                system=(
                    "You are a strict resume-to-job matcher. Only classify skills as "
                    "present if they are grounded in the resume text. Return JSON only."
                ),
                use_cache=False,
            )
            parsed = self._parse_json_object(response)
            return {
                "missing_skills": self._normalize_required_skills(
                    parsed.get("missing_skills")
                ),
                "present_skills": self._normalize_required_skills(
                    parsed.get("present_skills")
                ),
            }
        except Exception as exc:
            logger.warning("Gap identification fell back to heuristic matching: %s", exc)
            lowered_resume = resume_text.lower()
            present_skills = [
                skill for skill in normalized_skills if skill.lower() in lowered_resume
            ]
            missing_skills = [
                skill for skill in normalized_skills if skill not in present_skills
            ]
            return {
                "missing_skills": missing_skills,
                "present_skills": present_skills,
            }

    async def _estimate_improved_score(
        self,
        score_before: float,
        edit_result: dict[str, Any],
        gap_skills: list[str],
        strategy_name: str,
    ) -> float:
        prompt = (
            "Estimate the improved resume-to-job match score after the listed edits.\n"
            "Return only a decimal number between 0 and 1.\n\n"
            f"Score before: {score_before:.6f}\n"
            f"Strategy: {strategy_name}\n"
            f"Gap skills: {json.dumps(gap_skills, ensure_ascii=False)}\n"
            f"Changes made: {json.dumps(edit_result.get('changes', []), ensure_ascii=False)}"
        )

        try:
            response = await self.llm.complete(
                prompt=prompt,
                system=(
                    "You estimate resume match improvement conservatively and reply "
                    "with a decimal only."
                ),
                use_cache=False,
            )
            estimated = self._extract_first_float(response)
        except Exception as exc:
            logger.warning("Improved score estimation fell back to score_before: %s", exc)
            estimated = score_before

        return max(score_before, min(1.0, estimated))

    def _save_edited_resume(
        self,
        edit_result: dict[str, Any],
        path: Path,
        original_text: str,
    ) -> str:
        document = Document()
        document.add_heading("AutoApply Resume Edit Preview", level=0)
        document.add_paragraph(f"Strategy: {self.strategy_name}")

        document.add_heading("Changes", level=1)
        changes = self._normalize_string_list(edit_result.get("changes"))
        if changes:
            for change in changes:
                document.add_paragraph(change, style="List Bullet")
        else:
            document.add_paragraph("No explicit changes were returned.")

        edited_sections = edit_result.get("edited_sections")
        if isinstance(edited_sections, dict) and edited_sections:
            document.add_heading("Edited Sections", level=1)
            for section_name, content in edited_sections.items():
                document.add_heading(str(section_name), level=2)
                document.add_paragraph(str(content))

        new_summary = str(edit_result.get("new_summary") or "").strip()
        if new_summary:
            document.add_heading("Updated Summary", level=1)
            document.add_paragraph(new_summary)

        reordered_skills = self._normalize_required_skills(
            edit_result.get("reordered_skills")
        )
        if reordered_skills:
            document.add_heading("Reordered Skills", level=1)
            for skill in reordered_skills:
                document.add_paragraph(skill, style="List Bullet")

        document.add_heading("Original Resume", level=1)
        original_lines = [line.strip() for line in original_text.splitlines() if line.strip()]
        if original_lines:
            for line in original_lines:
                document.add_paragraph(line)
        else:
            document.add_paragraph("Original resume text was empty.")

        document.save(str(path))
        return str(path)

    def _build_flag_reason(self, match_score_after: float) -> str:
        return (
            f"Estimated match score {match_score_after:.3f} is below the required "
            f"threshold of {settings.resume_match_threshold:.3f}."
        )

    @staticmethod
    def _parse_json_object(value: str) -> dict[str, Any]:
        stripped = value.strip()
        fenced_match = re.search(r"```(?:json)?\s*(\{.*\})\s*```", stripped, re.DOTALL)
        if fenced_match:
            stripped = fenced_match.group(1).strip()
        elif not (stripped.startswith("{") and stripped.endswith("}")):
            json_match = re.search(r"(\{.*\})", stripped, re.DOTALL)
            if json_match:
                stripped = json_match.group(1).strip()

        parsed = json.loads(stripped)
        if not isinstance(parsed, dict):
            raise ValueError("Expected JSON object response")
        return parsed

    @staticmethod
    def _extract_first_float(value: str) -> float:
        match = re.search(r"[-+]?(?:\d+\.\d+|\d+|\.\d+)", value)
        if not match:
            raise ValueError("No numeric score found in response")
        return float(match.group(0))

    @staticmethod
    def _normalize_required_skills(value: Any) -> list[str]:
        normalized: list[str] = []
        seen: set[str] = set()

        def _add(item: Any) -> None:
            if item is None:
                return
            if isinstance(item, (list, tuple, set)):
                for nested in item:
                    _add(nested)
                return

            text = str(item).strip()
            if not text:
                return

            for part in re.split(r"[\n,]+", text):
                skill = part.strip()
                if not skill:
                    continue
                key = skill.lower()
                if key in seen:
                    continue
                seen.add(key)
                normalized.append(skill)

        _add(value)
        return normalized

    @staticmethod
    def _normalize_string_list(value: Any) -> list[str]:
        normalized: list[str] = []
        seen: set[str] = set()

        def _add(item: Any) -> None:
            if item is None:
                return
            if isinstance(item, (list, tuple, set)):
                for nested in item:
                    _add(nested)
                return

            text = str(item).strip()
            if not text:
                return

            lines = [line.strip() for line in text.splitlines() if line.strip()]
            if not lines:
                lines = [text]

            for line in lines:
                key = line.lower()
                if key in seen:
                    continue
                seen.add(key)
                normalized.append(line)

        _add(value)
        return normalized


__all__ = [
    "BASE_RESUME_PATH",
    "RESUME_DIR",
    "ResumeEditorAgent",
    "cosine_similarity",
]
