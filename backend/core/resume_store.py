"""
resume_store.py — Central helpers for resume file and user profile management.

Resume is stored at DATA_DIR/base_resume.docx.
User profile is stored at DATA_DIR/user_profile.json.
"""
from __future__ import annotations

import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
RESUME_PATH = DATA_DIR / "base_resume.docx"
PROFILE_PATH = DATA_DIR / "user_profile.json"

DATA_DIR.mkdir(parents=True, exist_ok=True)


# ---------------------------------------------------------------------------
# Resume helpers
# ---------------------------------------------------------------------------

def get_resume_metadata() -> dict[str, Any] | None:
    """Return metadata about the stored resume, or None if no resume exists."""
    if not RESUME_PATH.exists():
        return None
    stat = RESUME_PATH.stat()
    return {
        "uploaded": True,
        "filename": RESUME_PATH.name,
        "size_bytes": stat.st_size,
        "last_modified": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
    }


def get_resume_text() -> str:
    """
    Extract plain text from the stored base_resume.docx.
    Falls back to a minimal placeholder if no resume is uploaded yet.
    """
    if not RESUME_PATH.exists():
        logger.warning("No resume found at %s — using placeholder text", RESUME_PATH)
        return _placeholder_resume_text()

    try:
        from docx import Document  # type: ignore
        document = Document(str(RESUME_PATH))
        paragraphs = [
            p.text.strip()
            for p in document.paragraphs
            if p.text and p.text.strip()
        ]
        if paragraphs:
            return "\n".join(paragraphs)
        logger.warning("Resume DOCX parsed but contained no text — using placeholder")
    except Exception as exc:
        logger.exception("Failed to read resume DOCX: %s", exc)

    return _placeholder_resume_text()


def save_resume(file_bytes: bytes, original_filename: str) -> dict[str, Any]:
    """
    Accept raw bytes of a .docx or .pdf file, convert PDF→DOCX if needed,
    persist as base_resume.docx, and return metadata.
    """
    original_filename_lower = original_filename.lower()

    if original_filename_lower.endswith(".docx"):
        _write_resume_bytes(file_bytes)

    elif original_filename_lower.endswith(".pdf"):
        docx_bytes = _pdf_to_docx(file_bytes, original_filename)
        _write_resume_bytes(docx_bytes)

    else:
        raise ValueError(
            f"Unsupported resume file type: {original_filename}. "
            "Please upload a .docx or .pdf file."
        )

    metadata = get_resume_metadata()
    assert metadata is not None
    return metadata


def _write_resume_bytes(data: bytes) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    RESUME_PATH.write_bytes(data)
    logger.info("Resume saved to %s (%d bytes)", RESUME_PATH, len(data))


def _pdf_to_docx(pdf_bytes: bytes, original_filename: str) -> bytes:
    """Convert PDF bytes → DOCX bytes using PyMuPDF + python-docx."""
    import io
    import tempfile

    try:
        import fitz  # PyMuPDF  # type: ignore
        from docx import Document  # type: ignore

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp_pdf:
            tmp_pdf.write(pdf_bytes)
            tmp_pdf_path = tmp_pdf.name

        pdf_doc = fitz.open(tmp_pdf_path)
        full_text = "\n".join(
            pdf_doc.load_page(i).get_text("text")
            for i in range(pdf_doc.page_count)
        )
        pdf_doc.close()

        word_doc = Document()
        word_doc.add_heading(original_filename.replace(".pdf", ""), level=0)
        for line in full_text.splitlines():
            stripped = line.strip()
            if stripped:
                word_doc.add_paragraph(stripped)

        buf = io.BytesIO()
        word_doc.save(buf)
        buf.seek(0)
        return buf.read()

    except ImportError as exc:
        raise RuntimeError(
            "PyMuPDF (fitz) or python-docx is not installed — cannot convert PDF. "
            "Please upload a .docx file instead."
        ) from exc


def _placeholder_resume_text() -> str:
    return (
        "Apoorv Nath Tripathi\n"
        "apoorvtripathi99@gmail.com | +91 92051 02348 | linkedin.com/in/noobapoo99\n\n"
        "Education\n"
        "BITS Pilani, Pilani — B.E. (Hons.) | Nov 2021 – May 2025\n\n"
        "Experience\n"
        "Jr Software Engineer - Q3 Technologies | Jan 2026 - Present\n"
        "Contributing as Junior iOS Developer for FirstGroup (UK) internal staff app.\n"
        "Developing features for train delay monitoring, capacity tracking, duty assignment workflows.\n"
        "Built SwiftUI interfaces using view composition, custom modifiers, MVVM pattern.\n\n"
        "Full Stack Developer - Zenopsys.Ai | Sept 2025 - Nov 2025\n"
        "Built production-grade system using Golang, Next.js, REST APIs.\n"
        "Automated email notifications using AWS SES, AWS Lambda.\n\n"
        "Research Intern - National Chemical Laboratory, Pune | April 2024 – Dec 2024\n"
        "Applied machine learning, data preprocessing, statistical modeling; shortlisted for RACPI patent.\n\n"
        "Web Development Intern - Indian Red Cross Society | May 2023 – July 2023\n"
        "Designed and developed full-stack MERN blood donation management system.\n\n"
        "Projects\n"
        "LitedIn - MERN stack application with Prisma ORM, role-based access control.\n"
        "PersonalAgent - Agentic AI assistant using Python, LangGraph, LangChain, Groq, Ollama.\n\n"
        "Technical Skills\n"
        "Languages: C++, JavaScript, Python\n"
        "Frameworks: React.js, Next.js, Node.js, Express.js, Tailwind CSS, SwiftUI\n"
        "Databases & Tools: MongoDB, PostgreSQL, Git, GitHub, Postman\n"
        "Cloud & AI: AWS (Lambda, SES), FastAPI, LLMs, Groq, Ollama, LangChain, LangGraph\n"
        "Core CS: Data Structures and Algorithms, Object-Oriented Programming\n\n"
        "Achievements\n"
        "Top 10 in Here's Hackathon. Top 15 in Walmart Hackathon.\n"
        "Second round of Flipkart Grid 5.0 among 150,000 competitors.\n"
        "Solved 500+ LeetCode/GFG problems, top 350 at BITS Pilani.\n"
        "450+ open source contributions on GitHub.\n"
    )


# ---------------------------------------------------------------------------
# User profile helpers
# ---------------------------------------------------------------------------

def get_user_profile() -> dict[str, Any] | None:
    """Return saved user profile dict, or None if not set."""
    if not PROFILE_PATH.exists():
        return None
    try:
        return json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
    except Exception as exc:
        logger.exception("Failed to read user profile: %s", exc)
        return None


def save_user_profile(profile: dict[str, Any]) -> dict[str, Any]:
    """Persist the user profile and return it."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    PROFILE_PATH.write_text(
        json.dumps(profile, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    logger.info("User profile saved to %s", PROFILE_PATH)
    return profile


__all__ = [
    "DATA_DIR",
    "RESUME_PATH",
    "PROFILE_PATH",
    "get_resume_metadata",
    "get_resume_text",
    "save_resume",
    "get_user_profile",
    "save_user_profile",
]
